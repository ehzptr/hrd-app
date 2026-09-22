"""Reusable worksheet and template helpers for Excel workbooks."""

from typing import Optional

import polars as pl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from reporting.styles import BODY_FONT, HEADER_FILL, HEADER_FONT, LIGHT_GRAY, TEMPLATE_BORDER


def write_table(
    ws: Worksheet,
    df: Optional[pl.DataFrame],
    start_row: int = 1,
    start_col: int = 1,
) -> tuple[int, int]:
    """Write a styled Polars table and return its last row and column."""
    if df is None or df.is_empty():
        cell = ws.cell(row=start_row, column=start_col, value="(Tidak ada data)")
        cell.font = Font(name=BODY_FONT.name, size=BODY_FONT.sz, italic=True, color="7F7F7F")
        return start_row, start_col
    for j, col_name in enumerate(df.columns):
        cell = ws.cell(row=start_row, column=start_col + j, value=str(col_name))
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for i, row in enumerate(df.iter_rows(named=False), start=1):
        for j, value in enumerate(row):
            cell = ws.cell(row=start_row + i, column=start_col + j, value=value)
            if i % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=LIGHT_GRAY)
    last_row = start_row + len(df)
    last_col = start_col + len(df.columns) - 1
    ws.auto_filter.ref = (
        f"{get_column_letter(start_col)}{start_row}:{get_column_letter(last_col)}{last_row}"
    )
    ws.freeze_panes = ws.cell(row=start_row + 1, column=start_col).coordinate
    return last_row, last_col


def autofit(ws: Worksheet, max_scan_row: int = 2000) -> None:
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        longest = 0
        for row_idx in range(1, min(ws.max_row, max_scan_row) + 1):
            value = ws.cell(row_idx, col_idx).value
            if value is not None:
                longest = max(longest, len(str(value)))
        ws.column_dimensions[letter].width = min(max(longest + 2, 10), 42)


def data_sheet(wb: Workbook, name: str, df: Optional[pl.DataFrame]) -> Worksheet:
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    write_table(ws, df)
    autofit(ws)
    return ws


def style_template_sheet(
    ws: Worksheet,
    widths: dict[str, int],
    max_rows: int = 500,
) -> None:
    """Apply the common header/body treatment used by upload templates."""
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = TEMPLATE_BORDER
    for row in ws.iter_rows(min_row=2, max_row=max_rows, max_col=len(widths)):
        for cell in row:
            cell.font = BODY_FONT
            cell.border = TEMPLATE_BORDER
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{list(widths)[-1]}{max_rows}"


def style_guide_sheet(ws: Worksheet, title_row: int = 1) -> None:
    ws[f"A{title_row}"].font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    ws[f"A{title_row}"].fill = HEADER_FILL
    ws.merge_cells(start_row=title_row, start_column=1, end_row=title_row, end_column=2)
    for row in ws.iter_rows(min_row=title_row + 1):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")
