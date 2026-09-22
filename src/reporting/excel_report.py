"""
Comprehensive Excel report generator (Section 13).

Renders the complete audit package — from the executive dashboard with
KPI cards and native openpyxl charts down to the verbatim audit trail of
raw scans and rejected rows.
"""

from __future__ import annotations

import io
import hashlib
import json
import datetime as dt
from typing import Optional

import polars as pl
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from application.config import AppConfig, config_to_dataframe
from reporting.excel_utils import (
    autofit as _autofit,
    data_sheet as _data_sheet,
    write_table as _write_table,
)
from reporting.models import ReportData
from reporting.styles import (
    ACCENT_BLUE,
    BORDER_GRAY,
    CARD_BORDER,
    CARD_LABEL_FONT,
    CARD_VALUE_FONT,
    DARK_SLATE,
    HEADER_FILL,
    HEADER_FONT,
    INSIGHT_FONT,
    INSIGHT_HEADER_FONT,
    LIGHT_GRAY,
    NAVY,
    SECTION_FILL,
    SECTION_FONT,
    THIN_SIDE,
    WHITE,
)


# ---------------------------------------------------------------------------
# KPI card helper (used on the dashboard)
# ---------------------------------------------------------------------------


def _kpi_card(ws: Worksheet, row: int, col: int, label: str, value: str, accent: str = ACCENT_BLUE) -> None:
    label_cell = ws.cell(row=row, column=col, value=label)
    label_cell.font = CARD_LABEL_FONT
    label_cell.alignment = Alignment(horizontal="center", vertical="center")
    label_cell.fill = PatternFill("solid", fgColor=accent)
    label_cell.border = CARD_BORDER

    val_cell = ws.cell(row=row + 1, column=col, value=value)
    val_cell.font = CARD_VALUE_FONT
    val_cell.alignment = Alignment(horizontal="center", vertical="center")
    val_cell.border = CARD_BORDER


def _section_title(ws: Worksheet, row: int, col: int, text: str, span: int = 6) -> None:
    for c in range(col, col + span):
        cell = ws.cell(row=row, column=c)
        cell.fill = SECTION_FILL
    title_cell = ws.cell(row=row, column=col, value=text)
    title_cell.font = SECTION_FONT
    title_cell.alignment = Alignment(vertical="center")
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)


# ---------------------------------------------------------------------------
# Dashboard builder
# ---------------------------------------------------------------------------


def _build_dashboard(wb: Workbook, data: ReportData) -> None:
    ws = wb.create_sheet("DASHBOARD")
    ws.sheet_view.showGridLines = False

    # Title block
    ws.merge_cells("A1:K1")
    title = ws.cell(row=1, column=1, value="DEALERSHIP HR & SALES PERFORMANCE DASHBOARD")
    title.font = Font(name="Segoe UI", size=14, bold=True, color=WHITE)
    title.fill = HEADER_FILL
    title.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 28

    if data.report_period_label:
        ws.merge_cells("A2:K2")
        sub = ws.cell(row=2, column=1, value=f"Periode Laporan: {data.report_period_label}")
        sub.font = Font(name="Segoe UI", size=10, italic=True, color="595959")

    # KPI summary numbers
    summary = data.employee_summary
    total_emp = summary["No."].n_unique() if not summary.is_empty() else 0
    total_late = int(summary["Terlambat"].sum()) if not summary.is_empty() and "Terlambat" in summary.columns else 0
    total_absent = int(summary["Mangkir"].sum()) if not summary.is_empty() and "Mangkir" in summary.columns else 0
    avg_score = summary["Attendance_Score"].mean() if not summary.is_empty() and "Attendance_Score" in summary.columns else 100.0

    sales_perf = data.sales_performance
    total_spk = int(sales_perf["SPK"].sum()) if not sales_perf.is_empty() and "SPK" in sales_perf.columns else 0
    total_delivery = int(sales_perf["Delivery"].sum()) if not sales_perf.is_empty() and "Delivery" in sales_perf.columns else 0
    total_revenue = float(sales_perf["Revenue"].sum()) if not sales_perf.is_empty() and "Revenue" in sales_perf.columns else 0.0

    card_row = 4
    _kpi_card(ws, card_row, 1, "Total Karyawan", str(total_emp))
    _kpi_card(ws, card_row, 3, "Avg Attendance Score", f"{avg_score:.1f}%")
    _kpi_card(ws, card_row, 5, "Total Kejadian Telat", str(total_late))
    _kpi_card(ws, card_row, 7, "Total Kejadian Mangkir", str(total_absent))
    _kpi_card(ws, card_row, 9, "Total SPK Sales", str(total_spk))
    _kpi_card(ws, card_row, 11, "Total Delivery Sales", str(total_delivery))

    # Automated HR Insights section
    insight_row = 8
    _section_title(ws, insight_row, 1, "Rangkuman Insight Otomatis (HR & Sales Highlights)", span=11)
    ws.row_dimensions[insight_row].height = 20
    for idx, text in enumerate(data.insights, start=1):
        r = insight_row + idx
        c = ws.cell(row=r, column=1, value=f"\u2022 {text}")
        c.font = INSIGHT_FONT
        c.alignment = Alignment(vertical="center")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=11)

    # Tables: Watchlist (top anomalies) + Sales Leaderboard
    table_start_row = insight_row + len(data.insights) + 2

    # Watchlist: 5 lowest Attendance_Score
    _section_title(ws, table_start_row, 1, "Perhatian Khusus: 5 Karyawan Skor Terendah", span=5)
    watch_cols = ["No.", "Name", "Department", "Attendance_Score", "Terlambat", "Mangkir"]
    avail_watch = [c for c in watch_cols if c in summary.columns] if not summary.is_empty() else []
    watch = summary.sort("Attendance_Score").head(5).select(avail_watch) if avail_watch else pl.DataFrame()
    row = table_start_row + 1
    last_row, last_col = _write_table(ws, watch, start_row=row, start_col=1)

    # Top 5 Sales
    sales_col = 7
    _section_title(ws, table_start_row, sales_col, "Top Performa Sales", span=5)
    sales_cols = ["No.", "Name", "SPK", "Delivery", "Performance_Score"]
    avail_sales = [c for c in sales_cols if c in sales_perf.columns] if not sales_perf.is_empty() else []
    top_sales = sales_perf.sort("Performance_Score", descending=True).head(5).select(avail_sales) if avail_sales else pl.DataFrame()
    sales_row = table_start_row + 1
    _write_table(ws, top_sales, start_row=sales_row, start_col=sales_col)

    # Daily Trend Chart
    trend = data.daily_trend
    if not trend.is_empty():
        chart_section_row = max(last_row, sales_row + len(top_sales)) + 3
        _section_title(ws, chart_section_row, 1, "Trend Kehadiran Harian", span=11)
        data_start = chart_section_row + 1
        trend_cols = ["Tanggal", "Work_Days", "Present", "Late", "Absent", "Early_Leave"]
        avail_trend = [c for c in trend_cols if c in trend.columns]
        _write_table(ws, trend.select(avail_trend), start_row=data_start, start_col=1)

        # Line chart for Trend
        n = len(trend)
        chart = LineChart()
        chart.title = "Trend Hadir vs Terlambat vs Mangkir"
        chart.style = 13
        chart.y_axis.title = "Jumlah Orang"
        chart.x_axis.title = "Tanggal"
        chart.height = 12
        chart.width = 22

        cats = Reference(ws, min_col=1, min_row=data_start + 1, max_row=data_start + n)
        vals = Reference(ws, min_col=3, max_col=min(5, len(avail_trend)), min_row=data_start, max_row=data_start + n)
        chart.add_data(vals, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, f"G{data_start}")

        # Bar chart for Late vs Absent
        bar = BarChart()
        bar.type = "col"
        bar.title = "Perbandingan Terlambat vs Mangkir"
        bar.height = 8
        bar.width = 22
        vals2 = Reference(ws, min_col=4, max_col=5, min_row=data_start, max_row=data_start + n)
        bar.add_data(vals2, titles_from_data=True)
        bar.set_categories(cats)
        ws.add_chart(bar, f"G{data_start + 17}")

    dept_summary = data.department_summary
    if not dept_summary.is_empty():
        dept_chart_row = data_start + len(trend) + 35
        _section_title(ws, dept_chart_row, 1, "Attendance Score per Departemen")
        dept_data_row = dept_chart_row + 1
        dept_cols = ["Department", "Attendance_Score_%"]
        avail_dept = [c for c in dept_cols if c in dept_summary.columns]
        _write_table(
            ws,
            dept_summary.select(avail_dept),
            start_row=dept_data_row,
            start_col=1,
        )
        n = len(dept_summary)
        dept_bar = BarChart()
        dept_bar.type = "bar"
        dept_bar.title = "Attendance Score per Departemen"
        dept_bar.height = 8
        dept_bar.width = 22
        cats = Reference(ws, min_col=1, min_row=dept_data_row + 1, max_row=dept_data_row + n)
        vals = Reference(ws, min_col=2, min_row=dept_data_row, max_row=dept_data_row + n)
        dept_bar.add_data(vals, titles_from_data=True)
        dept_bar.set_categories(cats)
        ws.add_chart(dept_bar, f"G{dept_data_row}")

    _autofit(ws, max_scan_row=30)


# ---------------------------------------------------------------------------
# README / Metadata Sheet
# ---------------------------------------------------------------------------


README_ROWS = [
    ("LEMBAR DOKUMENTASI & METADATA LAPORAN", ""),
    ("Tujuan Laporan", "Rekapitulasi kehadiran karyawan, evaluasi kedisiplinan, analisis performa sales, dan payroll."),
    ("Format & Standar", "Dihasilkan secara otomatis oleh sistem HR Dashboard terstandar."),
    ("", ""),
    ("STRUKTUR SHEET", ""),
    ("DASHBOARD", "Ringkasan visual eksekutif: KPI utama, watchlist kedisiplinan, leaderboard sales, dan grafik."),
    ("CONFIG", "Snapshot konfigurasi aturan absensi, toleransi telat, bobot KPI sales, dan tarif potongan payroll."),
    ("EMPLOYEE_MASTER", "Master data karyawan aktif: ID, Nama, Divisi, Jabatan, Tipe (OFFICE/SALES), Gaji Pokok."),
    ("RAW_LOG", "Transaksi mentah dari mesin absensi setelah validasi dasar. Tidak diubah/dihapus (audit trail)."),
    ("DAILY_ATTENDANCE", "Hasil perhitungan harian per karyawan: scan pertama/terakhir, status, menit telat, lembur, anomali."),
    ("EMPLOYEE_SUMMARY", "KPI kehadiran per karyawan: attendance rate, punctuality, attendance score, tidak absen pulang, lembur."),
    ("ANOMALY", "Hanya baris dengan anomali (severity HIGH/LOW, lupa scan pulang, scan hari libur) untuk ditelusuri HR."),
    ("PAYROLL", "Estimasi potongan gaji: telat (per menit/kejadian), mangkir, dan potongan 1x tidak absen pulang khusus tim Sales."),
    ("SALES_PERFORMANCE", "KPI dan skor performa sales, dihitung HANYA dari Sales_Activity — bukan dari fingerprint."),
    ("Sales_Activity", "Input aktivitas sales (prospect, visit, test drive, SPK, delivery, revenue) per tanggal."),
    ("", ""),
    ("CATATAN PENTING", ""),
    (
        "Aturan First/Last Scan",
        "Scan pertama = absen masuk, scan terakhir = absen pulang. Scan di antaranya TIDAK ditafsirkan "
        "sebagai istirahat/keluar-masuk kantor, dicatat sebagai potensi anomali untuk audit.",
    ),
    (
        "Sales & Canvassing",
        "Mesin absensi tidak mencatat aktivitas canvassing sales. Jumlah scan TIDAK digunakan sebagai KPI sales. "
        "Performa sales murni berasal dari sheet Sales_Activity.",
    ),
    (
        "Data Ditolak",
        "Baris raw yang tidak valid (ID kosong / tanggal tidak terbaca) tidak dihitung tetapi tetap disimpan "
        "terpisah untuk audit (lihat sheet Data_Ditolak jika ada).",
    ),
]


def _timestamp(value: dt.datetime) -> str:
    if value.tzinfo is None:
        value = value.replace(tzinfo=dt.timezone.utc)
    return value.astimezone(dt.timezone.utc).isoformat()


def _build_readme(wb: Workbook, data: ReportData) -> None:
    ws = wb.create_sheet("README")
    ws.sheet_view.showGridLines = False
    rows = [
        *README_ROWS[:3],
        ("Generated At (UTC)", _timestamp(data.generated_at)),
        ("Report Period", data.report_period_label or "(not specified)"),
        ("Input Files", ", ".join(data.source_files) or "(not provided)"),
        ("", ""),
        *README_ROWS[3:],
    ]
    for i, (label, desc) in enumerate(rows, start=1):
        c1 = ws.cell(row=i, column=1, value=label)
        c2 = ws.cell(row=i, column=2, value=desc)
        if desc == "" and label:
            c1.font = Font(bold=True, size=13, color=WHITE)
            c1.fill = PatternFill("solid", fgColor=NAVY)
            ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)
        else:
            c1.font = Font(bold=True, color=NAVY)
            c2.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 110


def _build_provenance(wb: Workbook, data: ReportData) -> None:
    """Write machine-readable provenance and row-count metadata."""
    ws = wb.create_sheet("PROVENANCE")
    ws.sheet_view.showGridLines = False
    config_df = config_to_dataframe(data.config)
    row_counts = {
        "CONFIG": len(config_df),
        "EMPLOYEE_MASTER": len(data.employee_master),
        "RAW_LOG": len(data.raw_log),
        "DAILY_ATTENDANCE": len(data.daily_attendance),
        "EMPLOYEE_SUMMARY": len(data.employee_summary),
        "DEPARTMENT_SUMMARY": len(data.department_summary),
        "DAILY_TREND": len(data.daily_trend),
        "ANOMALY": len(data.anomalies),
        "PAYROLL": len(data.payroll),
        "SALES_PERFORMANCE": len(data.sales_performance),
        "Sales_Activity": len(data.sales_activity),
        "REJECTED": len(data.rejected),
    }
    config_hash = hashlib.sha256(
        json.dumps(config_df.to_dicts(), default=str, sort_keys=True).encode("utf-8")
    ).hexdigest()
    rows: list[tuple[str, object]] = [
        ("generated_at", _timestamp(data.generated_at)),
        ("report_period", data.report_period_label),
        ("source_files", json.dumps(data.source_files, ensure_ascii=False)),
        ("config_sha256", config_hash),
        ("generator", "hrd-app reporting"),
        ("sales_provenance", "Derived only from Sales_Activity, never fingerprint scans."),
    ]
    rows.extend((f"row_count.{name}", count) for name, count in row_counts.items())
    rows.extend((f"metadata.{key}", value) for key, value in data.metadata.items())
    for row, (label, value) in enumerate(rows, start=1):
        ws.cell(row=row, column=1, value=label).font = Font(bold=True, color=NAVY)
        ws.cell(row=row, column=2, value=value)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 100


# ---------------------------------------------------------------------------
# Public entrypoint
# ---------------------------------------------------------------------------


def generate_workbook(data: ReportData) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    _build_dashboard(wb, data)
    _data_sheet(wb, "CONFIG", config_to_dataframe(data.config))
    _data_sheet(wb, "EMPLOYEE_MASTER", data.employee_master)
    _data_sheet(wb, "RAW_LOG", data.raw_log)
    _data_sheet(wb, "DAILY_ATTENDANCE", data.daily_attendance)
    _data_sheet(wb, "EMPLOYEE_SUMMARY", data.employee_summary)
    _data_sheet(wb, "ANOMALY", data.anomalies)
    _data_sheet(wb, "PAYROLL", data.payroll)
    _data_sheet(wb, "SALES_PERFORMANCE", data.sales_performance)
    _data_sheet(wb, "Sales_Activity", data.sales_activity)
    if data.rejected is not None and not data.rejected.is_empty():
        _data_sheet(wb, "Data_Ditolak", data.rejected)
    _build_readme(wb, data)
    _build_provenance(wb, data)

    for ws in wb.worksheets:
        if ws.title != "DASHBOARD":
            ws.freeze_panes = ws.freeze_panes or "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()