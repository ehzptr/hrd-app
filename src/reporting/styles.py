"""Shared Excel style definitions used by reports and input templates."""

from openpyxl.styles import Border, Font, PatternFill, Side

NAVY = "1F4E78"
DARK_SLATE = "2F5597"
ACCENT_BLUE = "D9E1F2"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"
BORDER_GRAY = "D9D9D9"

HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(name="Segoe UI", size=10, bold=True, color=WHITE)
SECTION_FILL = PatternFill("solid", fgColor=DARK_SLATE)
SECTION_FONT = Font(name="Segoe UI", size=11, bold=True, color=WHITE)
CARD_VALUE_FONT = Font(name="Segoe UI", size=20, bold=True, color=NAVY)
CARD_LABEL_FONT = Font(name="Segoe UI", size=9, bold=False, color="595959")
INSIGHT_FONT = Font(name="Segoe UI", size=10, italic=False, color=NAVY)
INSIGHT_HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color=NAVY)
BODY_FONT = Font(name="Segoe UI", size=10, color="1F2937")
THIN_SIDE = Side(border_style="thin", color=BORDER_GRAY)
TEMPLATE_BORDER = Border(bottom=Side(style="thin", color="D1D5DB"))
CARD_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
