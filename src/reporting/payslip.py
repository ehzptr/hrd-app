"""
Payslip (Slip Gaji) generator for individual and bulk export in PDF and Excel formats.
Layout strictly designed after the professional Indonesian standard payslip specification.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from application.payroll import PayslipData

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, portrait
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import HRFlowable, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False


NAVY_HEX = "#1F4E78"
HEADER_BG_HEX = "#D9D9D9"
BORDER_GRAY_HEX = "#B0B0B0"
LIGHT_BG_HEX = "#F7F9FB"

NAVY_EXCEL = "1F4E78"
HEADER_BG_EXCEL = "D9D9D9"
BORDER_GRAY_EXCEL = "B0B0B0"
LIGHT_BG_EXCEL = "EBF1F5"

FONT_NAME = "Segoe UI"


def _fmt_currency(val: float) -> str:
    if val is None:
        return "0"
    return f"{val:,.0f}"


# ---------------------------------------------------------------------------
# PDF Generator (ReportLab)
# ---------------------------------------------------------------------------


def _build_payslip_flowables(slip: PayslipData, company_info: dict) -> list:
    """Build the flowable elements for a single payslip."""
    company_name = company_info.get("name", "PT. Dealership Maju Bersama")
    company_address = company_info.get("address", "Surabaya, Indonesia")

    styles = getSampleStyleSheet()

    # Title & Meta styles
    style_title = ParagraphStyle(
        "PayslipTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=22,
        leading=26,
        alignment=2,  # Right align
        textColor=colors.HexColor("#1A1A1A"),
    )
    style_company = ParagraphStyle(
        "CompanyHeader",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#1F4E78"),
    )
    style_address = ParagraphStyle(
        "CompanyAddress",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#595959"),
    )
    style_meta_label = ParagraphStyle(
        "MetaLabel",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#333333"),
    )
    style_meta_val = ParagraphStyle(
        "MetaValue",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#1A1A1A"),
    )
    style_sec_head = ParagraphStyle(
        "SectionHead",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=12,
        textColor=colors.HexColor("#1A1A1A"),
    )
    style_cell_left = ParagraphStyle(
        "CellLeft",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#1A1A1A"),
    )
    style_cell_right = ParagraphStyle(
        "CellRight",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        alignment=2,
        textColor=colors.HexColor("#1A1A1A"),
    )
    style_cell_bold_right = ParagraphStyle(
        "CellBoldRight",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=11,
        alignment=2,
        textColor=colors.HexColor("#1A1A1A"),
    )
    style_cell_bold_left = ParagraphStyle(
        "CellBoldLeft",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=11,
        textColor=colors.HexColor("#1A1A1A"),
    )

    flowables = []

    # 1. Header: Company Info on Left, "Slip Gaji" on Right
    header_data = [
        [
            Paragraph(f"<b>{company_name}</b><br/>{company_address}", style_company),
            Paragraph("<b>Slip Gaji</b>", style_title),
        ]
    ]
    t_header = Table(header_data, colWidths=[110 * mm, 70 * mm])
    t_header.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    flowables.append(t_header)
    flowables.append(Spacer(1, 3 * mm))

    # 2. Employee Metadata Bar (2 Columns of Key-Values)
    meta_data = [
        [
            Paragraph("Nama / NIK", style_meta_label),
            Paragraph(f": {slip.name} ({slip.employee_id})", style_meta_val),
            Paragraph("Tgl Mulai Bekerja", style_meta_label),
            Paragraph(f": {slip.join_date}", style_meta_val),
        ],
        [
            Paragraph("Dept / Jabatan", style_meta_label),
            Paragraph(f": {slip.department} / {slip.position}", style_meta_val),
            Paragraph("Periode Gaji", style_meta_label),
            Paragraph(f": {slip.period_label or '-'}", style_meta_val),
        ],
    ]
    t_meta = Table(meta_data, colWidths=[28 * mm, 62 * mm, 32 * mm, 58 * mm])
    t_meta.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 1),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
            ]
        )
    )
    flowables.append(t_meta)
    flowables.append(Spacer(1, 4 * mm))

    # 3. Financial Table (Pendapatan vs Potongan in 2 Columns)
    pendapatan_rows = [
        ("Gaji Pokok", _fmt_currency(slip.gaji_pokok)),
        ("Lembur", _fmt_currency(slip.lembur_pay)),
        ("Tunjangan Jabatan", _fmt_currency(slip.tunjangan_jabatan)),
        ("Uang Makan", _fmt_currency(slip.uang_makan)),
        ("Tunjangan Parkir / Transport", _fmt_currency(slip.tunjangan_transport)),
    ]
    potongan_rows = [
        ("Potongan Absen (Mangkir)", _fmt_currency(slip.potongan_mangkir)),
        ("Potongan Datang Terlambat", _fmt_currency(slip.potongan_telat)),
        ("Potongan Pulang Lebih Dulu", _fmt_currency(slip.potongan_pulang_cepat)),
        ("Potongan Lupa Absen Pulang (Sales)", _fmt_currency(slip.potongan_lupa_pulang_sales)),
        ("Potongan Lain-lain", _fmt_currency(slip.potongan_lain)),
    ]

    max_rows = max(len(pendapatan_rows), len(potongan_rows))

    fin_table_data = [
        [
            Paragraph("<b>Pendapatan</b>", style_sec_head),
            "",
            Paragraph("<b>Potongan</b>", style_sec_head),
            "",
        ]
    ]

    for i in range(max_rows):
        p_label, p_val = pendapatan_rows[i] if i < len(pendapatan_rows) else ("", "")
        d_label, d_val = potongan_rows[i] if i < len(potongan_rows) else ("", "")
        fin_table_data.append(
            [
                Paragraph(p_label, style_cell_left),
                Paragraph(p_val, style_cell_right),
                Paragraph(d_label, style_cell_left),
                Paragraph(d_val, style_cell_right),
            ]
        )

    # Total row
    fin_table_data.append(
        [
            Paragraph("<b>Total Pendapatan</b>", style_cell_bold_left),
            Paragraph(f"<b>{_fmt_currency(slip.total_pendapatan)}</b>", style_cell_bold_right),
            Paragraph("<b>Total Potongan</b>", style_cell_bold_left),
            Paragraph(f"<b>{_fmt_currency(slip.total_potongan)}</b>", style_cell_bold_right),
        ]
    )

    t_fin = Table(fin_table_data, colWidths=[55 * mm, 35 * mm, 55 * mm, 35 * mm])
    t_fin.setStyle(
        TableStyle(
            [
                # Header row styling (Gray fill)
                ("BACKGROUND", (0, 0), (1, 0), colors.HexColor(HEADER_BG_HEX)),
                ("BACKGROUND", (2, 0), (3, 0), colors.HexColor(HEADER_BG_HEX)),
                ("SPAN", (0, 0), (1, 0)),
                ("SPAN", (2, 0), (3, 0)),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                # Total row border
                ("LINEABOVE", (0, -1), (-1, -1), 1.2, colors.HexColor("#333333")),
                ("LINEBELOW", (0, -1), (-1, -1), 1.2, colors.HexColor("#333333")),
                ("BOTTOMPADDING", (0, -1), (-1, -1), 5),
                ("TOPPADDING", (0, -1), (-1, -1), 5),
            ]
        )
    )
    flowables.append(t_fin)
    flowables.append(Spacer(1, 3 * mm))

    # 4. Take Home Pay Banner
    thp_data = [
        [
            Paragraph("<b>Gaji Bersih / Take Home Pay</b>", style_sec_head),
            Paragraph(f"<b>Rp {_fmt_currency(slip.take_home_pay)}</b>", style_title),
        ]
    ]
    t_thp = Table(thp_data, colWidths=[90 * mm, 90 * mm])
    t_thp.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EBF1F5")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#1F4E78")),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    flowables.append(t_thp)
    flowables.append(Spacer(1, 4 * mm))

    # 5. Rangkuman Informasi Kehadiran
    att_head = [
        [
            Paragraph("<b>Rangkuman Informasi Kehadiran</b>", style_sec_head),
            "",
            "",
            "",
        ]
    ]
    att_rows = [
        ("Kehadiran", f"{slip.hari_kehadiran} Hari", "Terlambat", f"{slip.kali_terlambat} Kali ({slip.menit_terlambat} Menit)"),
        ("Ketidak Hadiran (Mangkir)", f"{slip.hari_mangkir} Hari", "Pulang Lebih Dulu", f"{slip.kali_pulang_cepat} Kali"),
        ("Cuti", f"{slip.hari_cuti} Hari", "Lupa Absen", f"{slip.kali_lupa_absen} Kali"),
        ("Izin", f"{slip.hari_izin} Hari", "Total Jam Lembur", f"{slip.total_jam_lembur:.1f} Jam"),
        ("Sakit", f"{slip.hari_sakit} Hari", "", ""),
    ]

    att_table_data = att_head
    for r in att_rows:
        att_table_data.append(
            [
                Paragraph(r[0], style_cell_left),
                Paragraph(r[1], style_cell_right),
                Paragraph(r[2], style_cell_left),
                Paragraph(r[3], style_cell_right),
            ]
        )

    t_att = Table(att_table_data, colWidths=[50 * mm, 40 * mm, 50 * mm, 40 * mm])
    t_att.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(HEADER_BG_HEX)),
                ("SPAN", (0, 0), (-1, 0)),
                ("TOPPADDING", (0, 0), (-1, -1), 2.5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LINEBELOW", (0, -1), (-1, -1), 0.5, colors.HexColor(BORDER_GRAY_HEX)),
            ]
        )
    )
    flowables.append(t_att)

    return flowables


def generate_payslip_pdf(slip: PayslipData, company_info: dict) -> bytes:
    """Generate a single employee payslip as a PDF document."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=portrait(A4),
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )
    flowables = _build_payslip_flowables(slip, company_info)
    doc.build(flowables)
    return buffer.getvalue()


def generate_bulk_payslip_pdf(slips: list[PayslipData], company_info: dict) -> bytes:
    """Generate a multi-page PDF document containing all employee payslips (1 slip per page)."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=portrait(A4),
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )
    all_flowables = []
    for idx, slip in enumerate(slips):
        slip_flowables = _build_payslip_flowables(slip, company_info)
        all_flowables.extend(slip_flowables)
        if idx < len(slips) - 1:
            all_flowables.append(PageBreak())

    doc.build(all_flowables)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Excel Generator (OpenPyXL)
# ---------------------------------------------------------------------------


def _populate_excel_payslip_sheet(ws, slip: PayslipData, company_info: dict) -> None:
    """Populate an openpyxl Worksheet with a beautifully formatted payslip."""
    company_name = company_info.get("name", "PT. Dealership Maju Bersama")
    company_address = company_info.get("address", "Surabaya, Indonesia")

    ws.sheet_view.showGridLines = False

    font_title = Font(name=FONT_NAME, size=18, bold=True, color="1A1A1A")
    font_company = Font(name=FONT_NAME, size=11, bold=True, color=NAVY_EXCEL)
    font_address = Font(name=FONT_NAME, size=9, color="595959")
    font_meta_lbl = Font(name=FONT_NAME, size=9, bold=True, color="333333")
    font_meta_val = Font(name=FONT_NAME, size=9, color="1A1A1A")
    font_head = Font(name=FONT_NAME, size=10, bold=True, color="1A1A1A")
    font_bold = Font(name=FONT_NAME, size=10, bold=True, color="1A1A1A")
    font_regular = Font(name=FONT_NAME, size=9, color="1A1A1A")
    font_thp_lbl = Font(name=FONT_NAME, size=11, bold=True, color=NAVY_EXCEL)
    font_thp_val = Font(name=FONT_NAME, size=14, bold=True, color=NAVY_EXCEL)

    fill_gray_head = PatternFill("solid", fgColor=HEADER_BG_EXCEL)
    fill_thp = PatternFill("solid", fgColor=LIGHT_BG_EXCEL)

    thin_border = Side(style="thin", color=BORDER_GRAY_EXCEL)
    border_full = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
    border_top_bottom = Border(top=thin_border, bottom=thin_border)

    # 1. Company Header & Title
    ws["A1"] = company_name
    ws["A1"].font = font_company
    ws["A2"] = company_address
    ws["A2"].font = font_address

    ws["D1"] = "Slip Gaji"
    ws["D1"].font = font_title
    ws["D1"].alignment = Alignment(horizontal="right")
    ws.merge_cells("D1:E1")

    # 2. Metadata
    ws["A4"] = "Nama / NIK"
    ws["A4"].font = font_meta_lbl
    ws["B4"] = f"{slip.name} ({slip.employee_id})"
    ws["B4"].font = font_meta_val

    ws["D4"] = "Tgl Mulai Bekerja"
    ws["D4"].font = font_meta_lbl
    ws["E4"] = slip.join_date
    ws["E4"].font = font_meta_val

    ws["A5"] = "Dept / Jabatan"
    ws["A5"].font = font_meta_lbl
    ws["B5"] = f"{slip.department} / {slip.position}"
    ws["B5"].font = font_meta_val

    ws["D5"] = "Periode Gaji"
    ws["D5"].font = font_meta_lbl
    ws["E5"] = slip.period_label or "-"
    ws["E5"].font = font_meta_val

    # 3. Financial Table Header
    ws["A7"] = "Pendapatan"
    ws["A7"].font = font_head
    ws["A7"].fill = fill_gray_head
    ws.merge_cells("A7:B7")

    ws["D7"] = "Potongan"
    ws["D7"].font = font_head
    ws["D7"].fill = fill_gray_head
    ws.merge_cells("D7:E7")

    # Financial Rows
    p_rows = [
        ("Gaji Pokok", slip.gaji_pokok),
        ("Lembur", slip.lembur_pay),
        ("Tunjangan Jabatan", slip.tunjangan_jabatan),
        ("Uang Makan", slip.uang_makan),
        ("Tunjangan Parkir / Transport", slip.tunjangan_transport),
    ]
    d_rows = [
        ("Potongan Absen (Mangkir)", slip.potongan_mangkir),
        ("Potongan Datang Terlambat", slip.potongan_telat),
        ("Potongan Pulang Lebih Dulu", slip.potongan_pulang_cepat),
        ("Potongan Lupa Absen Pulang (Sales)", slip.potongan_lupa_pulang_sales),
        ("Potongan Lain-lain", slip.potongan_lain),
    ]

    for i in range(max(len(p_rows), len(d_rows))):
        r_idx = 8 + i
        if i < len(p_rows):
            ws[f"A{r_idx}"] = p_rows[i][0]
            ws[f"A{r_idx}"].font = font_regular
            ws[f"B{r_idx}"] = p_rows[i][1]
            ws[f"B{r_idx}"].font = font_regular
            ws[f"B{r_idx}"].number_format = "#,##0"
        if i < len(d_rows):
            ws[f"D{r_idx}"] = d_rows[i][0]
            ws[f"D{r_idx}"].font = font_regular
            ws[f"E{r_idx}"] = d_rows[i][1]
            ws[f"E{r_idx}"].font = font_regular
            ws[f"E{r_idx}"].number_format = "#,##0"

    tot_row = 8 + max(len(p_rows), len(d_rows))
    ws[f"A{tot_row}"] = "Total Pendapatan"
    ws[f"A{tot_row}"].font = font_bold
    ws[f"B{tot_row}"] = slip.total_pendapatan
    ws[f"B{tot_row}"].font = font_bold
    ws[f"B{tot_row}"].number_format = "#,##0"

    ws[f"D{tot_row}"] = "Total Potongan"
    ws[f"D{tot_row}"].font = font_bold
    ws[f"E{tot_row}"] = slip.total_potongan
    ws[f"E{tot_row}"].font = font_bold
    ws[f"E{tot_row}"].number_format = "#,##0"

    for col in ["A", "B", "D", "E"]:
        ws[f"{col}{tot_row}"].border = border_top_bottom

    # 4. Take Home Pay Row
    thp_row = tot_row + 2
    ws[f"A{thp_row}"] = "Gaji Bersih / Take Home Pay"
    ws[f"A{thp_row}"].font = font_thp_lbl
    ws[f"A{thp_row}"].fill = fill_thp
    ws.merge_cells(f"A{thp_row}:C{thp_row}")

    ws[f"D{thp_row}"] = slip.take_home_pay
    ws[f"D{thp_row}"].font = font_thp_val
    ws[f"D{thp_row}"].fill = fill_thp
    ws[f"D{thp_row}"].number_format = '"Rp " #,##0'
    ws.merge_cells(f"D{thp_row}:E{thp_row}")

    # 5. Rangkuman Informasi Kehadiran
    att_head_row = thp_row + 2
    ws[f"A{att_head_row}"] = "Rangkuman Informasi Kehadiran"
    ws[f"A{att_head_row}"].font = font_head
    ws[f"A{att_head_row}"].fill = fill_gray_head
    ws.merge_cells(f"A{att_head_row}:E{att_head_row}")

    att_data = [
        ("Kehadiran", f"{slip.hari_kehadiran} Hari", "Terlambat", f"{slip.kali_terlambat} Kali ({slip.menit_terlambat} Menit)"),
        ("Ketidak Hadiran (Mangkir)", f"{slip.hari_mangkir} Hari", "Pulang Lebih Dulu", f"{slip.kali_pulang_cepat} Kali"),
        ("Cuti", f"{slip.hari_cuti} Hari", "Lupa Absen", f"{slip.kali_lupa_absen} Kali"),
        ("Izin", f"{slip.hari_izin} Hari", "Total Jam Lembur", f"{slip.total_jam_lembur:.1f} Jam"),
        ("Sakit", f"{slip.hari_sakit} Hari", "", ""),
    ]

    for idx, row in enumerate(att_data, start=att_head_row + 1):
        ws[f"A{idx}"] = row[0]
        ws[f"A{idx}"].font = font_regular
        ws[f"B{idx}"] = row[1]
        ws[f"B{idx}"].font = font_regular
        ws[f"B{idx}"].alignment = Alignment(horizontal="right")

        ws[f"D{idx}"] = row[2]
        ws[f"D{idx}"].font = font_regular
        ws[f"E{idx}"] = row[3]
        ws[f"E{idx}"].font = font_regular
        ws[f"E{idx}"].alignment = Alignment(horizontal="right")

    # Column Widths
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 4
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["E"].width = 20


def generate_payslip_excel(slip: PayslipData, company_info: dict) -> bytes:
    """Generate a single employee payslip as an Excel spreadsheet."""
    wb = Workbook()
    ws = wb.active
    ws.title = f"Slip_{slip.employee_id}"
    _populate_excel_payslip_sheet(ws, slip, company_info)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generate_bulk_payslip_excel(slips: list[PayslipData], company_info: dict) -> bytes:
    """Generate a consolidated Excel workbook with 1 sheet per employee payslip."""
    wb = Workbook()
    wb.remove(wb.active)  # Remove default active sheet

    for slip in slips:
        # Sheet title max 31 chars
        clean_name = "".join(c for c in slip.name if c.isalnum() or c in " _-")[:18].strip()
        sheet_title = f"{slip.employee_id}_{clean_name}"[:31]
        ws = wb.create_sheet(title=sheet_title)
        _populate_excel_payslip_sheet(ws, slip, company_info)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()