"""
Employee Master (Section 9).

Employee Type detection: if ``Department`` contains "SALES" the employee
defaults to type SALES, otherwise OFFICE — but HR can always override this
via the ``Employee_Type`` column in the uploaded master file, as required
by the spec.
"""

from __future__ import annotations

import datetime as dt
import io
from typing import Optional

import polars as pl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from reporting.excel_utils import style_guide_sheet, style_template_sheet

EMPLOYEE_MASTER_COLUMNS = [
    "No.",
    "Name",
    "Department",
    "Position",
    "Employee_Type",
    "Jam_Masuk",
    "Jam_Pulang",
    "Monthly_Salary",
    "Active",
]

EMPLOYEE_MASTER_SCHEMA = {
    "No.": pl.String,
    "Name": pl.String,
    "Department": pl.String,
    "Position": pl.String,
    "Employee_Type": pl.String,
    "Jam_Masuk": pl.String,
    "Jam_Pulang": pl.String,
    "Monthly_Salary": pl.Float64,
    "Active": pl.Boolean,
}

VALID_EMPLOYEE_TYPES = {"OFFICE", "SALES"}

NAVY = "1F4E78"
WHITE = "FFFFFF"


class EmployeeMasterError(ValueError):
    """Raised when the uploaded employee master file is malformed."""


def _clean_col(name: object) -> str:
    return str(name).replace("\n", " ").strip()


def _parse_time_str(val: object) -> Optional[str]:
    if val is None:
        return None
    if isinstance(val, (dt.time,)):
        return val.strftime("%H:%M")
    if isinstance(val, (dt.datetime,)):
        return val.strftime("%H:%M")
    s = str(val).strip()
    if not s or s.lower() in ("nan", "none", "nat", "<null>", ""):
        return None
    s = s.replace(".", ":")
    parts = s.split(":")
    if len(parts) >= 2:
        try:
            h = int(parts[0])
            m = int(parts[1])
            return f"{h:02d}:{m:02d}"
        except ValueError:
            return None
    return None


def detect_employee_type(department: Optional[str]) -> str:
    if department and "sales" in str(department).lower():
        return "SALES"
    return "OFFICE"


def load_employee_master(raw_input: Optional[pl.DataFrame]) -> pl.DataFrame:
    """Validate and normalize an uploaded Employee Master file.

    Returns a Polars DataFrame with exactly ``EMPLOYEE_MASTER_COLUMNS``, one row
    per unique Employee ID (last occurrence wins so HR can re-upload an
    updated file). Missing optional fields are filled with safe defaults;
    ``No.`` is mandatory.
    """
    if raw_input is None or raw_input.is_empty():
        return pl.DataFrame(schema=EMPLOYEE_MASTER_SCHEMA)

    df = raw_input
    rename_map = {c: _clean_col(c) for c in df.columns}
    df = df.rename(rename_map)

    id_col = next((c for c in ["No.", "Employee ID", "ID", "NIK"] if c in df.columns), None)
    if id_col is None:
        raise EmployeeMasterError(
            "Employee Master harus memiliki kolom employee ID (No./Employee ID/ID/NIK)."
        )

    # 1. No.
    s_no = df[id_col].cast(pl.String).str.strip_chars()

    # 2. Name
    if "Name" in df.columns:
        s_name = df["Name"].cast(pl.String).str.strip_chars()
    else:
        s_name = pl.Series("Name", [None] * len(df), dtype=pl.String)

    # 3. Department
    if "Department" in df.columns:
        s_dept = df["Department"].cast(pl.String).str.strip_chars()
    else:
        s_dept = pl.Series("Department", [None] * len(df), dtype=pl.String)

    # 4. Position
    if "Position" in df.columns:
        s_pos = df["Position"].cast(pl.String).str.strip_chars()
    else:
        s_pos = pl.Series("Position", [None] * len(df), dtype=pl.String)

    # 5. Employee Type
    if "Employee_Type" in df.columns:
        s_etype = df["Employee_Type"].cast(pl.String).str.strip_chars().str.to_uppercase()
    elif "Employee Type" in df.columns:
        s_etype = df["Employee Type"].cast(pl.String).str.strip_chars().str.to_uppercase()
    else:
        s_etype = pl.Series("Employee_Type", [None] * len(df), dtype=pl.String)

    # 6. Jam_Masuk & Jam_Pulang
    in_col = next(
        (c for c in ["Jam_Masuk", "Jam Masuk", "Jam_Kerja_Masuk", "Jam Kerja Masuk", "In", "Clock_In", "Clock In"] if c in df.columns),
        None,
    )
    out_col = next(
        (c for c in ["Jam_Pulang", "Jam Pulang", "Jam_Kerja_Pulang", "Jam Kerja Pulang", "Out", "Clock_Out", "Clock Out"] if c in df.columns),
        None,
    )
    range_col = next(
        (c for c in ["Jam_Kerja", "Jam Kerja", "Range_Jam_Kerja", "Schedule", "Jadwal"] if c in df.columns),
        None,
    )

    if in_col is not None:
        in_list = [_parse_time_str(v) for v in df[in_col].to_list()]
    elif range_col is not None:
        def _get_start(val):
            if val is None:
                return None
            txt = str(val)
            for sep in ["-", "s/d", "to"]:
                if sep in txt:
                    return _parse_time_str(txt.split(sep)[0])
            return _parse_time_str(txt)
        in_list = [_get_start(v) for v in df[range_col].to_list()]
    else:
        in_list = [None] * len(df)

    if out_col is not None:
        out_list = [_parse_time_str(v) for v in df[out_col].to_list()]
    elif range_col is not None:
        def _get_end(val):
            if val is None:
                return None
            txt = str(val)
            for sep in ["-", "s/d", "to"]:
                if sep in txt:
                    return _parse_time_str(txt.split(sep)[1])
            return None
        out_list = [_get_end(v) for v in df[range_col].to_list()]
    else:
        out_list = [None] * len(df)

    s_in = pl.Series("Jam_Masuk", in_list, dtype=pl.String)
    s_out = pl.Series("Jam_Pulang", out_list, dtype=pl.String)

    # 7. Monthly_Salary
    if "Monthly_Salary" in df.columns:
        s_salary = df["Monthly_Salary"].cast(pl.Float64, strict=False)
    elif "Monthly Salary" in df.columns:
        s_salary = df["Monthly Salary"].cast(pl.Float64, strict=False)
    else:
        s_salary = pl.Series("Monthly_Salary", [None] * len(df), dtype=pl.Float64)

    # 8. Active
    if "Active" in df.columns:
        s_act_raw = df["Active"].cast(pl.String).str.strip_chars().str.to_lowercase()
        s_active = ~s_act_raw.is_in(["n", "no", "tidak", "0", "false", "inactive"])
    else:
        s_active = pl.Series("Active", [True] * len(df), dtype=pl.Boolean)

    out = pl.DataFrame(
        [s_no.alias("No."), s_name.alias("Name"), s_dept.alias("Department"),
         s_pos.alias("Position"), s_etype.alias("Employee_Type"),
         s_in.alias("Jam_Masuk"), s_out.alias("Jam_Pulang"),
         s_salary.alias("Monthly_Salary"), s_active.alias("Active")]
    )

    out = out.filter(pl.col("No.").is_not_null() & (pl.col("No.") != ""))

    # Fallback Employee_Type from Department
    out = out.with_columns(
        Employee_Type=pl.when(
            pl.col("Employee_Type").is_in(list(VALID_EMPLOYEE_TYPES))
        )
        .then(pl.col("Employee_Type"))
        .otherwise(
            pl.when(pl.col("Department").str.to_lowercase().str.contains("sales"))
            .then(pl.lit("SALES"))
            .otherwise(pl.lit("OFFICE"))
        ),
        Name=pl.when(pl.col("Name").is_null() | (pl.col("Name") == ""))
        .then(pl.concat_str([pl.lit("Karyawan "), pl.col("No.")]))
        .otherwise(pl.col("Name")),
        Department=pl.when(pl.col("Department").is_null() | (pl.col("Department") == ""))
        .then(pl.lit("Belum Dipetakan"))
        .otherwise(pl.col("Department")),
        Position=pl.when(pl.col("Position").is_null()).then(pl.lit("")).otherwise(pl.col("Position")),
    )

    out = out.unique(subset=["No."], keep="last")
    return out.select(EMPLOYEE_MASTER_COLUMNS)


def merge_master_into_log(
    clean_log: pl.DataFrame, master: Optional[pl.DataFrame]
) -> tuple[pl.DataFrame, pl.DataFrame]:
    """Enrich attendance rows with employee master identity/type/schedule.

    Where an employee exists in the raw log but not in the master file, the
    identity fields already present on the log (or safe fallbacks) are kept
    and the employee is flagged in the returned "unmapped" DataFrame so HR
    can extend the master file. No attendance rows are dropped.
    """
    clean = clean_log
    if master is None or master.is_empty():
        enriched = clean.with_columns(
            Employee_Type=pl.when(pl.col("Department").str.to_lowercase().str.contains("sales"))
            .then(pl.lit("SALES"))
            .otherwise(pl.lit("OFFICE")),
            Position=pl.lit(""),
            Jam_Masuk_Master=pl.lit(None, dtype=pl.String),
            Jam_Pulang_Master=pl.lit(None, dtype=pl.String),
        )
        unmapped = enriched.select(["No.", "Name", "Department"]).unique()
        return enriched, unmapped

    m = master

    # Master columns to join
    m_cols = ["No.", "Name", "Department", "Employee_Type", "Position"]
    if "Jam_Masuk" in m.columns:
        m_cols.append("Jam_Masuk")
    if "Jam_Pulang" in m.columns:
        m_cols.append("Jam_Pulang")

    m_subset = m.select(m_cols)
    m_renamed = m_subset.rename(
        {
            "Name": "_m_Name",
            "Department": "_m_Department",
            "Employee_Type": "_m_Employee_Type",
            "Position": "_m_Position",
            **({"Jam_Masuk": "Jam_Masuk_Master"} if "Jam_Masuk" in m_cols else {}),
            **({"Jam_Pulang": "Jam_Pulang_Master"} if "Jam_Pulang" in m_cols else {}),
        }
    )

    joined = clean.join(m_renamed, on="No.", how="left")

    if "Jam_Masuk_Master" not in joined.columns:
        joined = joined.with_columns(Jam_Masuk_Master=pl.lit(None, dtype=pl.String))
    if "Jam_Pulang_Master" not in joined.columns:
        joined = joined.with_columns(Jam_Pulang_Master=pl.lit(None, dtype=pl.String))

    # Master identity wins for Name / Department if present
    enriched = joined.with_columns(
        Name=pl.coalesce([pl.col("_m_Name"), pl.col("Name")]),
        Department=pl.coalesce([pl.col("_m_Department"), pl.col("Department")]),
        Position=pl.coalesce([pl.col("_m_Position"), pl.lit("")]),
        Employee_Type=pl.coalesce(
            [
                pl.col("_m_Employee_Type"),
                pl.when(pl.col("Department").str.to_lowercase().str.contains("sales"))
                .then(pl.lit("SALES"))
                .otherwise(pl.lit("OFFICE")),
            ]
        ),
    ).drop(["_m_Name", "_m_Department", "_m_Employee_Type", "_m_Position"])

    master_ids = set(m["No."].to_list())
    unmapped = (
        enriched.filter(~pl.col("No.").is_in(master_ids))
        .select(["No.", "Name", "Department"])
        .unique()
    )
    return enriched, unmapped


def build_employee_master_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Employee Master"
    guide = wb.create_sheet("Petunjuk")

    ws.append(EMPLOYEE_MASTER_COLUMNS)
    ws.append(["EMP001", "Budi Santoso", "Sales Mobil Baru", "Sales Executive", "SALES", "08:30", "17:30", 5500000, "Y"])
    ws.append(["EMP002", "Sari Wijaya", "HRD & Admin", "Staff HRD", "OFFICE", "08:00", "17:00", 6000000, "Y"])

    style_template_sheet(
        ws,
        {"A": 14, "B": 22, "C": 20, "D": 20, "E": 14, "F": 14, "G": 14, "H": 16, "I": 10},
        max_rows=1000,
    )

    dv_type = DataValidation(type="list", formula1='"OFFICE,SALES"', allow_blank=True)
    dv_type.error = "Pilih OFFICE atau SALES."
    dv_type.errorTitle = "Employee Type tidak valid"
    ws.add_data_validation(dv_type)
    dv_type.add("E2:E1000")

    dv_active = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(dv_active)
    dv_active.add("I2:I1000")

    guide_rows = [
        ["PETUNJUK EMPLOYEE MASTER"],
        ["No.", "Employee ID — harus sama dengan kolom No. pada export mesin absensi."],
        ["Name", "Nama karyawan."],
        ["Department", "Departemen/divisi."],
        ["Position", "Jabatan (opsional)."],
        ["Employee_Type", "OFFICE atau SALES. Jika dikosongkan, sistem menebak dari Department (mengandung 'SALES')."],
        ["Jam_Masuk", "Jam masuk resmi per divisi/karyawan (format HH:MM, contoh: 08:30). Jika dikosongkan, mengikuti jadwal dashboard."],
        ["Jam_Pulang", "Jam pulang resmi per divisi/karyawan (format HH:MM, contoh: 17:30). Jika dikosongkan, mengikuti jadwal dashboard."],
        ["Monthly_Salary", "Gaji bulanan untuk perhitungan payroll (opsional, angka saja)."],
        ["Active", "Y = aktif, N = non-aktif. Karyawan non-aktif tetap dihitung di riwayat tapi ditandai."],
    ]
    for row in guide_rows:
        guide.append(row)
    style_guide_sheet(guide)
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 90

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()